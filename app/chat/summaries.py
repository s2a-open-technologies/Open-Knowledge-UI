# Open Knowledge UI - A privacy-focused, self-hosted RAG system
# SPDX-FileCopyrightText: 2026 S2A Open Technologies gGmbH
# SPDX-License-Identifier: AGPL-3.0-or-later
# Copyright (C) 2026 S2A Open Technologies gGmbH
#
# This program is free software: you can redistribute it and/or modify
# it under the terms of the GNU Affero General Public License as published by
# the Free Software Foundation, either version 3 of the License, or
# (at your option) any later version.
#
# This program is distributed in the hope that it will be useful,
# but WITHOUT ANY WARRANTY; without even the implied warranty of
# MERCHANTABILITY or FITNESS FOR A PARTICULAR PURPOSE. See the
# GNU Affero General Public License for more details.
#
# You should have received a copy of the GNU Affero General Public License
# along with this program. If not, see <https://www.gnu.org/licenses/>.

"""
Analysen (Summaries) gehoeren dem Nutzer, der sie erstellt: jeder hat einen eigenen Ordner
``<SUMMARIES_DIR>/u<id>/``. Weder die Oberflaeche noch die Werkzeuge des Modells sehen oder
aendern Dateien anderer Nutzer.
"""

import os
import shutil
from logging import getLogger
from pathlib import Path

logger = getLogger(__name__)


def base_dir() -> Path:
    return Path(os.getenv("SUMMARIES_DIR") or "./data/summaries")


def user_dir(user_id: int | None, create: bool = False) -> Path:
    """Ordner eines Nutzers; ohne Nutzer (Skripte) der gemeinsame Basisordner."""
    path = base_dir() if user_id is None else base_dir() / f"u{int(user_id)}"
    if create:
        path.mkdir(parents=True, exist_ok=True)
    return path


def safe_path(user_id: int | None, filename: str) -> Path | None:
    """Pfad einer .md-Datei im Ordner des Nutzers oder ``None`` (Ausbruchsversuch, falsche Endung)."""
    folder = user_dir(user_id).resolve()
    path = (folder / filename).resolve()
    if path.parent != folder or path.suffix != ".md":
        return None
    return path


def adopt_legacy(user_id: int) -> int:
    """Uebergibt Dateien aus der Zeit vor den Nutzerordnern (direkt im Basisordner) einem Admin."""
    base = base_dir()
    if not base.is_dir():
        return 0
    moved = 0
    for file in base.glob("*.md"):
        target = user_dir(user_id, create=True) / file.name
        if not target.exists():
            shutil.move(str(file), str(target))
            moved += 1
    if moved:
        logger.info(
            f"{moved} bestehende Analysen in den Ordner von Nutzer {user_id} verschoben"
        )
    return moved


def markdown_to_docx(content: str) -> bytes:
    """Markdown als DOCX, formatiert wie die Vorschau (Tabellen, Code-Bloecke, Listen, Links,
    Durchstreichen). markdown-it-py folgt derselben CommonMark/GFM-Spec wie marked in der
    Vorschau, damit das DOCX optisch der Markdown-Vorschau entspricht."""
    import io

    from bs4 import BeautifulSoup, NavigableString, Tag
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.shared import Pt, RGBColor
    from markdown_it import MarkdownIt

    # Markdown -> HTML via markdown-it-py (gfm-commonmark + table + strikethrough),
    # dieselbe Spec wie marked in der EasyMDE-Vorschau.
    md = MarkdownIt("commonmark", {"linkify": False, "html": False}).enable(
        ["table", "strikethrough"]
    )
    # GFM-Zeilenumbrüche: einzelner \n wird <br> (wie marked breaks:true).
    md.options["breaks"] = True
    html_body = md.render(content)

    # HTML -> DOCX via python-docx + BeautifulSoup
    doc = Document()

    # Basis-Stil etwas an Vorschau anpassen
    style = doc.styles["Normal"]
    style.font.size = Pt(11)
    style.font.name = "Calibri"

    soup = BeautifulSoup(html_body, "html.parser")

    def _add_runs_from_inline(paragraph, node, bold=False, italic=False, strike=False):
        """Rekursiv Inline-Knoten (strong/em/del/code/a/br/span/text) in Runs
        umwandeln, sodass Verschachtelung wie **_fett+kursiv_** erhalten bleibt."""
        if isinstance(node, NavigableString):
            text = str(node)
            if text:
                run = paragraph.add_run(text)
                run.bold = bold or None
                run.italic = italic or None
                run.font.strike = strike or None
            return
        if not isinstance(node, Tag):
            return
        name = node.name
        if name == "br":
            paragraph.add_run().add_break()
            return
        if name in ("strong", "b"):
            bold = True
        elif name in ("em", "i"):
            italic = True
        elif name in ("del", "s", "strike"):
            strike = True
        elif name == "code":
            for child in node.children:
                if isinstance(child, NavigableString):
                    run = paragraph.add_run(str(child))
                    run.bold = bold or None
                    run.italic = italic or None
                    run.font.strike = strike or None
                    run.font.name = "Consolas"
                    run.font.size = Pt(10)
                else:
                    _add_runs_from_inline(paragraph, child, bold, italic, strike)
            return
        elif name == "a":
            text = node.get_text()
            run = paragraph.add_run(text)
            run.bold = bold or None
            run.italic = italic or None
            run.font.strike = strike or None
            run.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
            run.underline = True
            return
        for child in node.children:
            _add_runs_from_inline(paragraph, child, bold, italic, strike)

    def _render_list(list_el, doc, depth=0):
        """Rendert eine <ul>/<ol> inkl. verschachtelter Listen."""
        is_ordered = list_el.name == "ol"
        for li in list_el.find_all("li", recursive=False):
            style_name = "List Number" if is_ordered else "List Bullet"
            para = doc.add_paragraph(style=style_name)
            for child in li.children:
                if isinstance(child, Tag) and child.name in ("ul", "ol"):
                    _render_list(child, doc, depth + 1)
                else:
                    _add_runs_from_inline(para, child)

    def _render_block(element, doc):
        """Rendert einen Block-Level-Knoten rekursiv."""
        if isinstance(element, NavigableString):
            if str(element).strip():
                para = doc.add_paragraph()
                _add_runs_from_inline(para, element)
            return
        if not isinstance(element, Tag):
            return

        if element.name in ("h1", "h2", "h3", "h4", "h5", "h6"):
            level = int(element.name[1])
            para = doc.add_heading(level=level)
            _add_runs_from_inline(para, element)

        elif element.name == "p":
            para = doc.add_paragraph()
            _add_runs_from_inline(para, element)

        elif element.name in ("ul", "ol"):
            _render_list(element, doc)

        elif element.name == "pre":
            code = element.find("code")
            text = code.get_text() if code else element.get_text()
            para = doc.add_paragraph()
            run = para.add_run(text.rstrip("\n"))
            run.font.name = "Consolas"
            run.font.size = Pt(10)
            # Leichte Hintergrundschattierung über Absatz-Eigenschaften ist
            # in python-docx aufwendig; wir belassen es bei der Schriftart.

        elif element.name == "blockquote":
            para = doc.add_paragraph(style="Intense Quote")
            for child in element.children:
                if (
                    isinstance(child, Tag)
                    and child.name == "p"
                    or isinstance(child, NavigableString)
                    and str(child).strip()
                ):
                    _add_runs_from_inline(para, child)
            # Weitere Block-Typen innerhalb von blockquote (Listen etc.)
            # rekursiv anhängen.
            for child in element.children:
                if isinstance(child, Tag) and child.name not in ("p",):
                    _render_block(child, doc)

        elif element.name == "hr":
            doc.add_page_break()

        elif element.name == "table":
            rows = element.find_all("tr", recursive=False)
            if not rows:
                # markdown-it packt <thead>/<tbody> um die <tr>
                thead = element.find("thead")
                tbody = element.find("tbody")
                rows = []
                if thead:
                    rows += thead.find_all("tr", recursive=False)
                if tbody:
                    rows += tbody.find_all("tr", recursive=False)
            if not rows:
                return
            num_cols = max(
                len(tr.find_all(["th", "td"], recursive=False)) for tr in rows
            )
            table = doc.add_table(rows=0, cols=num_cols, style="Table Grid")
            for tr in rows:
                row = table.add_row()
                cells = tr.find_all(["th", "td"], recursive=False)
                for cell_idx, cell in enumerate(cells):
                    if cell_idx >= num_cols:
                        break
                    cell_el = row.cells[cell_idx]
                    # Zelleninhalt als eigener Absatz mit Inline-Formatierung
                    cell_el.text = ""
                    para = cell_el.paragraphs[0]
                    _add_runs_from_inline(para, cell)
                    if cell.name == "th":
                        for paragraph in cell_el.paragraphs:
                            for run in paragraph.runs:
                                run.bold = True
            table.alignment = WD_TABLE_ALIGNMENT.CENTER

        elif element.name in ("div", "section", "article", "main", "thead", "tbody"):
            # Block-Container: Kinder rekursiv rendern (vorher wurde Inhalt
            # verworfen, was zu fehlenden Abschnitten im DOCX führte).
            for child in element.children:
                _render_block(child, doc)

        else:
            para = doc.add_paragraph()
            _add_runs_from_inline(para, element)

    for element in soup.children:
        _render_block(element, doc)

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
