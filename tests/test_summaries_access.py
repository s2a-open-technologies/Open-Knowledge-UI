from pathlib import Path

import pytest
from flask import Flask

from app.agent.tools.custom_tools import CustomTools
from app.api import register_blueprints
from app.chat import summaries
from app.rag.retrieval import RetrievalScope, set_scope
from tests.test_access import add_user, client_for


@pytest.fixture
def app(db):
    app = Flask(__name__)
    app.secret_key = "t"
    app.extensions["ingestion_status"] = {"is_ingesting": False, "knowledge_id": None}
    register_blueprints(app)
    return app


def test_users_cannot_read_change_or_delete_each_others_summaries(app, db):
    alice, bob = add_user(db, "alice", "user"), add_user(db, "bob", "user")
    a, b = client_for(app, alice), client_for(app, bob)

    assert (
        a.put("/api/summaries/geheim.md", json={"content": "nur Alice"}).status_code
        == 200
    )
    assert [s["filename"] for s in a.get("/api/summaries").json] == ["geheim.md"]

    assert b.get("/api/summaries").json == []
    assert b.get("/api/summaries/geheim.md").status_code == 404
    b.put(
        "/api/summaries/geheim.md", json={"content": "von Bob"}
    )  # legt eine EIGENE Datei an
    assert a.get("/api/summaries/geheim.md").json["content"] == "nur Alice"
    assert b.delete("/api/summaries/geheim.md").status_code == 200
    assert a.get("/api/summaries/geheim.md").json["content"] == "nur Alice"


def test_path_traversal_and_foreign_folders_are_rejected(app, db):
    alice, bob = add_user(db, "alice", "user"), add_user(db, "bob", "user")
    client_for(app, alice).put("/api/summaries/a.md", json={"content": "x"})
    b = client_for(app, bob)
    for name in (f"../u{alice['id']}/a.md", "../../app.db.md", "%2e%2e/x.md"):
        assert b.get(f"/api/summaries/{name}").status_code in (400, 404)
        assert b.put(f"/api/summaries/{name}", json={"content": "x"}).status_code in (
            400,
            404,
        )
    assert not any(Path(summaries.base_dir()).glob("*.md"))


def test_legacy_files_go_to_the_first_admin_only(app, db):
    base = summaries.base_dir()
    base.mkdir(parents=True, exist_ok=True)
    (base / "alt.md").write_text("Altbestand", encoding="utf-8")
    bob = add_user(db, "bob", "editor")
    assert (
        client_for(app, bob).get("/api/summaries").json == []
    )  # kein Admin: sieht nichts
    admin = {"id": 1, "role": "admin", "username": "admin"}
    assert [
        s["filename"] for s in client_for(app, admin).get("/api/summaries").json
    ] == ["alt.md"]


def test_model_tools_use_the_folder_of_the_requesting_user(db):
    tools = CustomTools(llm=None, retriever=None)
    set_scope(RetrievalScope((), user_id=7))
    tools.write_summary("meins.md", "Inhalt von 7")
    assert (summaries.base_dir() / "u7" / "meins.md").exists()
    set_scope(RetrievalScope((), user_id=8))
    assert "meins.md" not in str(tools.list_summaries())
    assert (
        "FEHLER" in tools.read_summary("meins.md")
        or "nicht" in tools.read_summary("meins.md").lower()
    )
    set_scope(None)


def test_deleting_a_summary_never_touches_another_users_tracking_row(app, db):
    alice, bob = add_user(db, "alice", "user"), add_user(db, "bob", "user")
    db.create_summary_file("gleich.md", alice["id"])
    db.create_summary_file("gleich.md", bob["id"])
    client_for(app, alice).put("/api/summaries/gleich.md", json={"content": "A"})
    client_for(app, bob).put("/api/summaries/gleich.md", json={"content": "B"})

    assert client_for(app, bob).delete("/api/summaries/gleich.md").status_code == 200
    remaining = [(s.filename, s.user_id) for s in db.get_all_summaries()]
    assert remaining == [("gleich.md", alice["id"])]  # Alices Eintrag bleibt
    assert client_for(app, alice).get("/api/summaries/gleich.md").json["content"] == "A"


def test_normal_users_download_their_own_summaries_as_md_and_docx(app, db):
    import io

    from docx import Document

    alice, bob = add_user(db, "alice", "user"), add_user(db, "bob", "user")
    a, b = client_for(app, alice), client_for(app, bob)
    a.put(
        "/api/summaries/bericht.md",
        json={"content": "# Titel\n\n| A | B |\n|---|---|\n| 1 | 2 |\n\n- **fett**"},
    )

    md = a.get("/api/summaries/bericht.md/download-md")
    assert md.status_code == 200
    assert "attachment" in md.headers["Content-Disposition"]
    assert md.data.decode().startswith("# Titel")

    docx = a.get("/api/summaries/bericht.md/download-docx")
    assert docx.status_code == 200
    assert "bericht.docx" in docx.headers["Content-Disposition"]
    doc = Document(io.BytesIO(docx.data))
    assert doc.paragraphs[0].text == "Titel"
    assert doc.tables[0].cell(1, 1).text == "2"

    for kind in ("md", "docx"):
        assert b.get(f"/api/summaries/bericht.md/download-{kind}").status_code == 404
        assert b.get(
            f"/api/summaries/../u{alice['id']}/bericht.md/download-{kind}"
        ).status_code in (400, 404)
